"""Exports generated from the current SQLite data: findings (CSV/XLSX/JSON) and an HTML project report.

Every export is written to EXPORT_DIRECTORY (kept for the Reports page) and returned for download.
Values from scanner reports are untrusted: CSV cells cannot start a formula, XLSX cells are stored
as text, and the HTML report escapes everything it prints.
"""

import csv
import html
import io
import json
import logging
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Any, Callable

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import __version__
from app.core.config import Settings
from app.core.errors import AppError, NotFoundError
from app.models import Finding, FindingStatus, Scan, ScanStatus, Severity, User
from app.services import audit_service, dashboard_service, remediation_service
from app.services.audit_service import AuditAction
from app.services.finding_service import FindingFilters, FindingSortField, SortOrder, all_findings
from app.services.project_service import get_project
from app.services.scan_history_service import SCAN_TIME

logger = logging.getLogger(__name__)

NOT_PROVIDED = "Not Provided"
NOT_EXTRACTED = "Not Extracted"
_SAFE_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,200}$")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


class ExportFormat(str, Enum):
    CSV = "csv"
    XLSX = "xlsx"
    JSON = "json"


_MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
    "html": "text/html; charset=utf-8",
}
_NOT_IN_REPORT = '<span class="missing">Not in report</span>'
_NOT_SET = '<span class="missing">Not set</span>'


@dataclass
class ExportFile:
    path: Path
    filename: str
    media_type: str
    count: int


def _iso(value: datetime | None) -> str | None:
    return value.isoformat() if value else None


# Column header -> value. Headers deliberately match the importer's synonyms, so an export re-imports.
COLUMNS: list[tuple[str, str, Callable[[Finding, int], Any]]] = [
    ("ID", "id", lambda f, n: f.id),
    ("Project", "project", lambda f, n: f.project.name),
    ("Title", "title", lambda f, n: f.normalized_title),
    ("Original Title", "original_title", lambda f, n: f.original_title),
    ("Severity", "severity", lambda f, n: f.normalized_severity.value if f.normalized_severity else None),
    ("Original Severity", "original_severity", lambda f, n: f.original_severity),
    ("Status", "status", lambda f, n: f.status.value),
    ("Assigned To", "assigned_to", lambda f, n: f.assignee.name if f.assignee else None),
    ("Category", "category", lambda f, n: f.category),
    ("CWE", "cwe", lambda f, n: f.cwe),
    ("OWASP", "owasp", lambda f, n: f.owasp),
    ("CVSS Score", "cvss_score", lambda f, n: f.cvss_score),
    ("Source Tool", "source_tool", lambda f, n: f.source_tool),
    ("Template / Plugin", "source_template", lambda f, n: f.source_template),
    ("Affected URL", "affected_url", lambda f, n: f.affected_url),
    ("Endpoint", "endpoint", lambda f, n: f.endpoint),
    ("HTTP Method", "http_method", lambda f, n: f.http_method),
    ("Parameter", "parameter", lambda f, n: f.parameter),
    ("Description", "description", lambda f, n: f.description),
    ("Impact", "impact", lambda f, n: f.impact),
    ("Evidence", "evidence", lambda f, n: f.evidence),
    ("Steps to Reproduce", "steps_to_reproduce", lambda f, n: f.steps_to_reproduce),
    ("Recommendation", "recommendation", lambda f, n: f.recommendation),
    ("References", "references", lambda f, n: f.source_reference),
    ("First Detected", "first_detected", lambda f, n: _iso(f.first_detected)),
    ("Last Detected", "last_detected", lambda f, n: _iso(f.last_detected)),
    ("Observations", "observations", lambda f, n: n),
    ("Fingerprint", "fingerprint", lambda f, n: f.finding_key),
]


def _csv_safe(value: Any) -> str:
    """Neutralise spreadsheet formula injection (OWASP CSV injection guidance)."""
    if value is None:
        return ""
    text = str(value)
    return "'" + text if text.startswith(_FORMULA_START) else text


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S-%f")


def _slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-").lower()[:60] or "project"


def _write(settings: Settings, name: str, data: bytes) -> Path:
    directory = settings.export_path
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / name
    with open(path, "xb") as handle:
        handle.write(data)
    return path


def _describe(filters: FindingFilters) -> dict[str, Any]:
    described = {
        "search": filters.search,
        "project_ids": filters.project_ids,
        "scan_id": filters.scan_id,
        "severities": filters.severities,
        "statuses": [s.value for s in filters.statuses],
        "source_tools": filters.source_tools,
        "categories": filters.categories,
        "assignee": filters.assignee,
        "date_field": filters.date_field.value if (filters.date_from or filters.date_to) else None,
        "date_from": filters.date_from.isoformat() if filters.date_from else None,
        "date_to": filters.date_to.isoformat() if filters.date_to else None,
        "remediation": filters.remediation,
    }
    return {key: value for key, value in described.items() if value not in (None, [], "")}


def export_findings(
    db: Session,
    settings: Settings,
    filters: FindingFilters,
    fmt: ExportFormat,
    actor: User | None,
    sort_by: FindingSortField = FindingSortField.SEVERITY,
    sort_order: SortOrder = SortOrder.DESC,
) -> ExportFile:
    rows = all_findings(db, filters, sort_by, sort_order)
    exported_at = datetime.now(timezone.utc)
    headers = [header for header, _, _ in COLUMNS]

    if fmt == ExportFormat.CSV:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        for finding, count in rows:
            writer.writerow([_csv_safe(get(finding, count)) for _, _, get in COLUMNS])
        data = buffer.getvalue().encode("utf-8-sig")  # BOM so Excel detects UTF-8
    elif fmt == ExportFormat.XLSX:
        data = _xlsx(headers, rows, filters, exported_at)
    else:
        payload = {
            "exported_at": exported_at.isoformat(),
            "application": f"VAPT Management Dashboard {__version__}",
            "filters": _describe(filters),
            "count": len(rows),
            "findings": [{key: get(finding, count) for _, key, get in COLUMNS} for finding, count in rows],
        }
        data = json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")

    name = f"{_stamp()}_findings.{fmt.value}"
    path = _write(settings, name, data)
    audit_service.record(
        db, actor, AuditAction.REPORT_EXPORTED, "export", None,
        {"kind": "findings", "format": fmt.value, "file": name, "count": len(rows), "filters": _describe(filters)},
        project_id=filters.project_ids[0] if len(filters.project_ids) == 1 else None,
    )
    db.commit()
    logger.info("Findings exported format=%s count=%d file=%s", fmt.value, len(rows), name)
    return ExportFile(path=path, filename=name, media_type=_MEDIA_TYPES[fmt.value], count=len(rows))


def _xlsx(headers: list[str], rows: list[tuple[Finding, int]], filters: FindingFilters, exported_at: datetime) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font

    # Write-only mode streams rows instead of keeping every cell object in memory.
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Findings")
    sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEFGHIJKLMNOPQRSTUVWXYZ", (8, 18, 45, 45, 14, 16, 14, 18, 16, 12, 14, 10, 14, 14, 45, 40, 10, 16)):
        sheet.column_dimensions[column].width = width
    bold, top = Font(bold=True), Alignment(wrap_text=False, vertical="top")

    def text_cell(value, **style):
        cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", value))  # control characters XLSX cannot store
        cell.data_type = "s"  # text, never a formula
        for key, val in style.items():
            setattr(cell, key, val)
        return cell

    sheet.append([text_cell(header, font=bold) for header in headers])
    for finding, count in rows:
        values = []
        for _, _, get in COLUMNS:
            value = get(finding, count)
            values.append(text_cell(value, alignment=top) if isinstance(value, str) else value)
        sheet.append(values)

    about = workbook.create_sheet("About")
    about.append(["Exported at", exported_at.isoformat()])
    about.append(["Findings", len(rows)])
    about.append(["Filters", json.dumps(_describe(filters)) or "none"])
    about.append(["Source", "Generated from the VAPT Management Dashboard database at export time."])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- Project report (HTML) -----------------------------------------------------------------

_SEVERITY_ORDER = [Severity.CRITICAL, Severity.HIGH, Severity.MEDIUM, Severity.LOW, Severity.INFORMATIONAL]
_SEVERITY_COLOURS = {"Critical": "#b91c1c", "High": "#ea580c", "Medium": "#d97706", "Low": "#0284c7", "Informational": "#64748b", "Unclassified": "#94a3b8"}


def _e(value: Any) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def _field(finding: Finding, name: str, value: Any) -> str:
    if value not in (None, ""):
        return _e(value)
    missing = (finding.source_metadata or {}).get("not_extracted") or []
    label = NOT_EXTRACTED if name in missing else NOT_PROVIDED
    return f'<span class="missing">{label}</span>'


def _fmt_date(value: datetime | None) -> str:
    return value.strftime("%d %b %Y %H:%M UTC") if value else "—"


def project_report(db: Session, settings: Settings, project_id: int, actor: User | None) -> ExportFile:
    project = get_project(db, project_id)
    summary = dashboard_service.get_summary(db, project_id)
    findings = all_findings(db, FindingFilters(project_ids=[project_id]))
    scans = db.scalars(select(Scan).where(Scan.project_id == project_id).order_by(SCAN_TIME.desc(), Scan.id.desc())).all()
    matrix_rows = db.execute(
        select(Finding.normalized_severity, Finding.status, func.count())
        .where(Finding.project_id == project_id)
        .group_by(Finding.normalized_severity, Finding.status)
    ).all()
    assignees = db.execute(
        select(User.name, func.count())
        .join(Finding, Finding.assigned_to == User.id)
        .where(Finding.project_id == project_id, Finding.status.in_(dashboard_service.ACTIVE_STATUSES))
        .group_by(User.name)
        .order_by(func.count().desc(), User.name)
    ).all()
    generated = datetime.now(timezone.utc)
    guidance = _guidance_groups(db, findings)
    page = _render_report(project, summary, findings, scans, matrix_rows, assignees, generated, guidance)

    name = f"{_stamp()}_project-report_{_slug(project.name)}.html"
    path = _write(settings, name, page.encode("utf-8"))
    audit_service.record(
        db, actor, AuditAction.REPORT_EXPORTED, "project", project.id,
        {"kind": "project_report", "format": "html", "file": name, "findings": len(findings)},
    )
    db.commit()
    logger.info("Project report exported project=%s findings=%d file=%s", project.id, len(findings), name)
    return ExportFile(path=path, filename=name, media_type=_MEDIA_TYPES["html"], count=len(findings))


def _guidance_groups(db: Session, findings: list[tuple[Finding, int]]) -> dict:
    """Findings per standard guideline (exact rule links only), and how many have none."""
    index = remediation_service.LibraryIndex.load(db)
    groups: dict[str, dict] = {}
    unlinked = 0
    for finding, _ in findings:
        matches = remediation_service.match_finding(finding, index)
        if not matches:
            unlinked += 1
        for match in matches:
            group = groups.setdefault(match.guideline.slug, {"guideline": match.guideline, "total": 0, "active": 0, "bases": set()})
            group["total"] += 1
            group["active"] += finding.status in dashboard_service.ACTIVE_STATUSES
            group["bases"].add("scanner rule" if match.match_type == "scanner_rule" else "CWE")
    ordered = sorted(groups.values(), key=lambda g: (-g["active"], -g["total"], g["guideline"].title))
    return {"groups": ordered, "unlinked": unlinked}


def _guidance_section(guidance: dict) -> str:
    if not guidance["groups"]:
        return '<p class="missing">No finding in this project matches a standard guideline by scanner rule or CWE.</p>'
    rows = []
    for group in guidance["groups"]:
        guideline = group["guideline"]
        links = "".join(
            '<br><a href="' + _e(ref["url"]) + '" rel="noopener noreferrer">' + _e(ref["title"]) + "</a>"
            for ref in guideline.references[:2]
        )
        rows.append(
            "<tr><td><strong>" + _e(guideline.title) + '</strong><br><span class="meta">' + _e(guideline.summary) + "</span>"
            + links + "</td><td>" + _e(guideline.category) + "</td><td>" + str(group["active"]) + "</td><td>"
            + str(group["total"]) + "</td><td>" + _e(", ".join(sorted(group["bases"]))) + "</td></tr>"
        )
    return (
        "<table><thead><tr><th>Standard guidance</th><th>Category</th><th>Active findings</th><th>All findings</th>"
        "<th>Linked by</th></tr></thead><tbody>" + "".join(rows) + "</tbody></table>"
        + '<p class="meta">' + str(guidance["unlinked"]) + " finding(s) have no standard guideline (no exact scanner rule or CWE match).</p>"
    )


def _render_report(project, summary, findings, scans, matrix_rows, assignees, generated, guidance) -> str:
    sev = summary.findings_by_severity
    counts = {
        "Critical": sev.critical, "High": sev.high, "Medium": sev.medium, "Low": sev.low,
        "Informational": sev.informational, "Unclassified": sev.unclassified,
    }
    statuses = {row.status.value: row.count for row in summary.status_summary}
    indicators = summary.indicators
    completed = [s for s in scans if s.processing_status == ScanStatus.COMPLETED]
    total = summary.finding_count
    fixed = statuses.get("Resolved", 0) + statuses.get("Closed", 0)
    exceptions = statuses.get("Accepted Risk", 0) + statuses.get("False Positive", 0)

    # Executive summary: factual statements computed from the database only.
    facts = []
    app_bits = ", ".join(_e(x) for x in (project.application_name, project.environment.value if project.environment else None) if x)
    facts.append(f"This report covers project <strong>{_e(project.name)}</strong>{f' ({app_bits})' if app_bits else ''}.")
    if completed:
        times = [s.scan_date or s.uploaded_at for s in completed]
        latest = completed[0]
        facts.append(
            f"It is based on {len(completed)} successfully processed scan(s) dated {_fmt_date(min(times))} to "
            f"{_fmt_date(max(times))}; the latest is “{_e(latest.scan_name)}” ({_e(latest.scanner.value if latest.scanner else '')})."
        )
    else:
        facts.append("No scans have been processed for this project.")
    if total:
        breakdown = ", ".join(f"{n} {label}" for label, n in counts.items() if n)
        facts.append(f"{total} unique finding(s) are recorded: {breakdown}.")
        facts.append(
            f"{indicators.active_findings} are active (Open, In Progress or Reopened), including "
            f"{indicators.active_critical_high} Critical or High. {fixed} are Resolved or Closed and {exceptions} are "
            f"Accepted Risk or False Positive."
        )
        if indicators.oldest_active_critical_high:
            facts.append(f"The oldest active Critical/High finding was first detected on {_fmt_date(indicators.oldest_active_critical_high)}.")
    else:
        facts.append("No findings have been recorded for this project.")

    matrix: dict[str, dict[str, int]] = {}
    for severity, status, count in matrix_rows:
        matrix.setdefault(severity.value if severity else "Unclassified", {})[status.value] = count
    status_names = [s.value for s in FindingStatus]
    severity_names = [s.value for s in _SEVERITY_ORDER] + (["Unclassified"] if counts["Unclassified"] else [])
    risk_table = "".join(
        f"<tr><th scope=\"row\">{_e(name)}</th>"
        + "".join(f"<td>{matrix.get(name, {}).get(st, 0) or ''}</td>" for st in status_names)
        + f"<td><strong>{counts.get(name, 0)}</strong></td></tr>"
        for name in severity_names
    )
    max_count = max(counts.values()) or 1
    distribution = "".join(
        f'<div class="bar-row"><span class="bar-label">{_e(name)}</span>'
        f'<span class="bar" style="width:{(counts[name] / max_count) * 100:.1f}%;background:{_SEVERITY_COLOURS[name]}"></span>'
        f'<span class="bar-value">{counts[name]}</span></div>'
        for name in severity_names
    )

    def severity_of(f: Finding) -> str:
        return f.normalized_severity.value if f.normalized_severity else "Unclassified"

    register_rows = "".join(
        f"<tr><td>{i}</td><td>{_e(f.normalized_title)}</td><td>{_e(severity_of(f))}</td><td>{_e(f.status.value)}</td>"
        f"<td>{_e(f.endpoint or '')}</td><td>{_e(f.cwe or '')}</td><td>{_e(f.source_tool)}</td><td>{_e(_fmt_date(f.last_detected))}</td></tr>"
        for i, (f, _) in enumerate(findings, start=1)
    )
    by_severity_sections = []
    for name in ("Critical", "High", "Medium", "Low"):
        group = [(f, n) for f, n in findings if severity_of(f) == name]
        body = (
            "".join(
                f"<tr><td>{_e(f.normalized_title)}</td><td>{_e(f.status.value)}</td><td>{_e(f.endpoint or '')}</td>"
                f"<td>{_e(f.assignee.name if f.assignee else 'Unassigned')}</td><td>{_e(_fmt_date(f.first_detected))}</td></tr>"
                for f, _ in group
            )
            if group
            else f'<tr><td colspan="5" class="missing">No {name} findings.</td></tr>'
        )
        by_severity_sections.append(
            f'<section><h3>{name} findings ({len(group)})</h3><table><thead><tr><th>Title</th><th>Status</th>'
            f"<th>Endpoint</th><th>Assigned to</th><th>First detected</th></tr></thead><tbody>{body}</tbody></table></section>"
        )

    def detail(index: int, f: Finding, observations: int) -> str:
        rows = [
            ("Severity", f"{_e(severity_of(f))} (report: {_field(f, 'original_severity', f.original_severity)})"),
            ("Status", _e(f.status.value)),
            ("Assigned to", _e(f.assignee.name if f.assignee else "Unassigned")),
            ("Source", f"{_e(f.source_tool)}{' · ' + _e(f.source_template) if f.source_template else ''}"),
            ("Affected URL", _field(f, "affected_url", f.affected_url)),
            ("Endpoint", _field(f, "endpoint", f.endpoint)),
            ("HTTP method", _field(f, "http_method", f.http_method)),
            ("Parameter", _field(f, "parameter", f.parameter)),
            ("CWE / OWASP / CVSS", f"{_field(f, 'cwe', f.cwe)} · {_field(f, 'owasp', f.owasp)} · {_field(f, 'cvss_score', f.cvss_score)}"),
            ("Detected", f"first {_e(_fmt_date(f.first_detected))}, last {_e(_fmt_date(f.last_detected))}, {observations} observation(s)"),
        ]
        blocks = [
            ("Description", f.description, "description", False),
            ("Impact", f.impact, "impact", False),
            ("Evidence", f.evidence, "evidence", True),
            ("Steps to reproduce", f.steps_to_reproduce, "steps_to_reproduce", True),
            ("Original recommendation (from the report)", f.recommendation, "recommendation", False),
            ("References", f.source_reference, "source_reference", False),
        ]
        facts_html = "".join(f"<tr><th scope=\"row\">{label}</th><td>{value}</td></tr>" for label, value in rows)
        blocks_html = "".join(
            f"<h5>{label}</h5>" + (f"<pre>{_e(value)}</pre>" if pre and value else f"<p>{_field(f, name, value)}</p>")
            for label, value, name, pre in blocks
        )
        return (
            f'<article class="finding"><h4>{index}. {_e(f.normalized_title)}</h4>'
            f'<table class="facts"><tbody>{facts_html}</tbody></table>{blocks_html}</article>'
        )

    details = "".join(detail(i, f, n) for i, (f, n) in enumerate(findings, start=1)) or '<p class="missing">No findings.</p>'
    assignee_rows = "".join(f"<tr><td>{_e(name)}</td><td>{count}</td></tr>" for name, count in assignees) or (
        '<tr><td colspan="2" class="missing">No active findings are assigned.</td></tr>'
    )
    scan_rows = "".join(
        f"<tr><td>{_e(s.scan_name)}</td><td>{_e(s.scanner.value if s.scanner else 'Not identified')}"
        f"{' ' + _e(s.report_type.value) if s.report_type else ''}</td><td>{_e(_fmt_date(s.scan_date)) if s.scan_date else _NOT_IN_REPORT}</td>"
        f"<td>{_e(_fmt_date(s.uploaded_at))}</td><td>{_e(s.processing_status.value)}</td><td>{s.total_findings}</td>"
        f"<td>{_e(s.original_filename)}</td></tr>"
        for s in scans
    ) or '<tr><td colspan="7" class="missing">No scans.</td></tr>'

    info_rows = [
        ("Project", project.name), ("Application", project.application_name), ("Environment", project.environment.value if project.environment else None),
        ("Application URL", project.application_url), ("Owner", project.owner), ("Status", project.status.value),
        ("Description", project.description),
    ]
    info = "".join(f"<tr><th scope=\"row\">{label}</th><td>{_e(value) if value else _NOT_SET}</td></tr>" for label, value in info_rows)

    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>VAPT project report – {_e(project.name)}</title>
<style>
body{{font-family:system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;color:#0f172a;margin:2rem auto;max-width:1100px;padding:0 1.25rem;line-height:1.45;font-size:14px}}
h1{{font-size:1.6rem;margin:0}} h2{{font-size:1.2rem;border-bottom:2px solid #0f766e;padding-bottom:.25rem;margin-top:2rem}}
h3{{font-size:1rem;margin-top:1.25rem}} h4{{font-size:1rem;margin:0 0 .5rem}} h5{{font-size:.8rem;text-transform:uppercase;color:#475569;margin:.75rem 0 .2rem}}
table{{border-collapse:collapse;width:100%;margin:.5rem 0}} th,td{{border:1px solid #cbd5e1;padding:.35rem .5rem;text-align:left;vertical-align:top}}
thead th{{background:#f1f5f9}} .facts th{{width:12rem;background:#f8fafc}} .meta{{color:#475569}}
pre{{white-space:pre-wrap;word-break:break-word;background:#f8fafc;border:1px solid #e2e8f0;padding:.5rem;font-size:12px}}
.missing{{color:#94a3b8;font-style:italic}} .finding{{border:1px solid #e2e8f0;border-radius:6px;padding:1rem;margin:1rem 0;break-inside:avoid-page}}
.bar-row{{display:flex;align-items:center;gap:.5rem;margin:.25rem 0}} .bar-label{{width:8rem}} .bar{{height:12px;border-radius:2px;min-width:2px}} .bar-value{{font-weight:600}}
@media print{{body{{margin:0;max-width:none}} h2{{break-after:avoid}}}}
</style></head><body>
<header><h1>VAPT project report</h1><p class="meta">{_e(project.name)} · generated {_e(_fmt_date(generated))} by VAPT Management Dashboard {__version__}.
All figures are calculated from the application database at the time of export.</p></header>

<h2>Project information</h2><table class="facts"><tbody>{info}</tbody></table>

<h2>Executive summary</h2>{"".join(f"<p>{fact}</p>" for fact in facts)}

<h2>Risk summary</h2>
<p>Unique findings by severity and workflow status.</p>
<table><thead><tr><th>Severity</th>{"".join(f"<th>{_e(s)}</th>" for s in status_names)}<th>Total</th></tr></thead><tbody>{risk_table}</tbody></table>

<h2>Severity distribution</h2>{distribution}

<h2>Findings by severity</h2>{"".join(by_severity_sections)}

<h2>Findings register</h2>
<table><thead><tr><th>#</th><th>Title</th><th>Severity</th><th>Status</th><th>Endpoint</th><th>CWE</th><th>Source Tool</th><th>Last Detected</th></tr></thead>
<tbody>{register_rows or '<tr><td colspan="8" class="missing">No findings.</td></tr>'}</tbody></table>

<h2>Detailed findings</h2>{details}

<h2>Remediation summary</h2>
<table class="facts"><tbody>
<tr><th scope="row">Needs action</th><td>{indicators.active_findings} (Open {statuses.get('Open', 0)}, In Progress {statuses.get('In Progress', 0)}, Reopened {statuses.get('Reopened', 0)})</td></tr>
<tr><th scope="row">Fixed</th><td>{fixed} (Resolved {statuses.get('Resolved', 0)}, Closed {statuses.get('Closed', 0)})</td></tr>
<tr><th scope="row">Exceptions</th><td>{exceptions} (Accepted Risk {statuses.get('Accepted Risk', 0)}, False Positive {statuses.get('False Positive', 0)})</td></tr>
<tr><th scope="row">Active and unassigned</th><td>{indicators.unassigned_active}</td></tr>
<tr><th scope="row">Resolved or closed in the last 30 days</th><td>{indicators.resolved_last_30_days}</td></tr>
</tbody></table>
<h3>Active findings by assignee</h3><table><thead><tr><th>Assignee</th><th>Active findings</th></tr></thead><tbody>{assignee_rows}</tbody></table>
<h3>Standard remediation guidance</h3>
<p class="meta">The recommendations in the detailed findings are the scanners' own text and are unchanged. The guidance below is
application guidance from the remediation library. It is linked to findings only by an exact scanner rule (ZAP alert ID,
Nuclei template) or a specific CWE, and is not part of the uploaded reports.</p>
{_guidance_section(guidance)}

<h2>Scan information</h2>
<table><thead><tr><th>Scan</th><th>Tool / format</th><th>Scan date</th><th>Uploaded</th><th>Processing</th><th>Findings</th><th>Original file</th></tr></thead><tbody>{scan_rows}</tbody></table>
</body></html>
"""


# --- Saved exports -------------------------------------------------------------------------

MAX_LISTED_EXPORTS = 200


def list_exports(settings: Settings) -> list[dict[str, Any]]:
    """The newest saved exports (older files stay on disk and remain downloadable by name)."""
    directory = settings.export_path
    if not directory.is_dir():
        return []
    files = []
    for path in directory.iterdir():
        if path.is_file() and _SAFE_NAME.match(path.name):
            stat = path.stat()
            kind = "project_report" if "_project-report_" in path.name else "findings" if "_findings." in path.name else "other"
            files.append({
                "name": path.name,
                "kind": kind,
                "format": path.suffix.lstrip(".").lower(),
                "size": stat.st_size,
                "created_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc),
            })
    files.sort(key=lambda f: f["name"], reverse=True)  # names start with a UTC timestamp
    return files[:MAX_LISTED_EXPORTS]


class ExportNotFoundError(NotFoundError):
    pass


def saved_export(settings: Settings, name: str) -> ExportFile:
    if not _SAFE_NAME.match(name):
        raise AppError("Invalid export file name.")
    root = settings.export_path.resolve()
    path = (root / name).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise ExportNotFoundError(f'Export "{name}" was not found.')
    media = _MEDIA_TYPES.get(path.suffix.lstrip(".").lower(), "application/octet-stream")
    return ExportFile(path=path, filename=name, media_type=media, count=0)
