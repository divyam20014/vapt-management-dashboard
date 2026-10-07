"""CSV and XLSX findings sheets with flexible column mapping (see parsers/tabular.py)."""

import csv
import io
import zipfile
from typing import Any

from app.models.enums import ReportType, Scanner
from app.parsers.base import BaseReportParser, ParsedReport, ReportFile, ReportParseError
from app.parsers.tabular import describe_columns, find_header, mapping_warnings, normalize_tabular_record, rows_to_records

MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000
MAX_COMPRESSION_RATIO = 200
csv.field_size_limit(16 * 1024 * 1024)


def _no_header_error(kind: str, rows: list[list[Any]]) -> ReportParseError:
    return ReportParseError(
        f"No findings table was recognised in this {kind}: none of the first rows has a title column "
        f"(e.g. \"Title\", \"Finding\", \"Vulnerability\") plus another known column such as \"Severity\". "
        f"Columns found: {describe_columns(rows)}. No findings were imported."
    )


class CsvParser(BaseReportParser):
    scanner = Scanner.IMPORTED
    report_type = ReportType.CSV
    extensions = (".csv",)

    def can_parse(self, file: ReportFile) -> bool:
        return True  # the only CSV format; problems are reported by parse()

    def _decode(self, file: ReportFile, warnings: list[str]) -> str:
        try:
            return file.content.decode("utf-8-sig")
        except UnicodeDecodeError:
            warnings.append("The file is not UTF-8; it was read as Windows-1252 (the encoding Excel uses for CSV on Windows).")
            return file.content.decode("cp1252", errors="strict")

    def parse(self, file: ReportFile) -> ParsedReport:
        warnings: list[str] = []
        try:
            text = self._decode(file, warnings)
        except UnicodeDecodeError as exc:
            raise ReportParseError("The CSV file is neither UTF-8 nor Windows-1252 text.") from exc
        # Only the delimiter is detected; quoting follows RFC 4180 / Excel ("" escapes a quote). Letting the
        # sniffer guess quoting rules mis-splits real exports that contain quotes and markup.
        try:
            delimiter = csv.Sniffer().sniff(text[:16384], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
        if delimiter != ",":
            warnings.append(f"Columns are separated by {'a tab' if delimiter == chr(9) else repr(delimiter)} rather than commas.")
        try:
            rows = [row for row in csv.reader(io.StringIO(text), csv.excel, delimiter=delimiter)]
        except csv.Error as exc:
            raise ReportParseError(f"The CSV file is malformed ({exc}).") from exc
        found = find_header(rows)
        if found is None:
            raise _no_header_error("CSV file", rows)
        header_index, mapping = found
        records = rows_to_records(rows[header_index + 1:], mapping, header_index + 2, {})
        warnings += mapping_warnings(mapping)
        if header_index:
            warnings.append(f"The first {header_index} row(s) above the column headings were not imported.")
        return ParsedReport(records=records, warnings=warnings)

    def normalize(self, record):
        return normalize_tabular_record(record)


class XlsxParser(BaseReportParser):
    scanner = Scanner.IMPORTED
    report_type = ReportType.XLSX
    extensions = (".xlsx",)

    def can_parse(self, file: ReportFile) -> bool:
        return True

    @staticmethod
    def _check_archive(content: bytes) -> None:
        try:
            archive = zipfile.ZipFile(io.BytesIO(content))
        except zipfile.BadZipFile as exc:
            raise ReportParseError("The file is not a valid XLSX workbook (it is not a readable ZIP package).") from exc
        entries = archive.infolist()
        names = {entry.filename for entry in entries}
        if "xl/workbook.xml" not in names:
            raise ReportParseError("The file is not an Excel XLSX workbook.")
        if any(name.lower().endswith("vbaproject.bin") for name in names):
            raise ReportParseError("Workbooks containing macros are not accepted.")
        total = sum(entry.file_size for entry in entries)
        compressed = sum(entry.compress_size for entry in entries) or 1
        if len(entries) > MAX_ZIP_ENTRIES or total > MAX_UNCOMPRESSED_BYTES or total / compressed > MAX_COMPRESSION_RATIO:
            raise ReportParseError("The workbook expands to an unsafe size and was not opened.")

    def parse(self, file: ReportFile) -> ParsedReport:
        from openpyxl import load_workbook  # local import keeps start-up light

        self._check_archive(file.content)
        try:
            workbook = load_workbook(io.BytesIO(file.content), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for damaged files
            raise ReportParseError("The XLSX workbook could not be read; it may be damaged.") from exc

        candidates = []
        first_rows: list[list[Any]] = []
        try:
            for sheet in workbook.worksheets:
                rows = [list(row) for row in sheet.iter_rows(values_only=True)]
                first_rows = first_rows or rows
                found = find_header(rows)
                if found:
                    candidates.append((sheet.title, rows, *found))
        finally:
            workbook.close()
        if not candidates:
            raise _no_header_error("workbook", first_rows)

        candidates.sort(key=lambda c: len(c[3].columns), reverse=True)
        if len(candidates) > 1 and len(candidates[0][3].columns) == len(candidates[1][3].columns):
            names = ", ".join(f'"{c[0]}"' for c in candidates)
            raise ReportParseError(
                f"Several sheets look like findings tables ({names}) and none is clearly the main one. "
                "Upload a workbook with a single findings sheet."
            )
        sheet_name, rows, header_index, mapping = candidates[0]
        records = rows_to_records(rows[header_index + 1:], mapping, header_index + 2, {"sheet": sheet_name})
        warnings = mapping_warnings(mapping, f'sheet "{sheet_name}"')
        if len(candidates) > 1:
            ignored = ", ".join(f'"{c[0]}"' for c in candidates[1:])
            warnings.append(f'Findings were read from sheet "{sheet_name}"; other sheets with fewer recognised columns were ignored: {ignored}.')
        return ParsedReport(records=records, warnings=warnings)

    def normalize(self, record):
        return normalize_tabular_record(record)
