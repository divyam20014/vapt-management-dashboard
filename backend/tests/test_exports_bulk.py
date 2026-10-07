"""Phase 5: findings exports, project report and bulk upload, checked against SQLite."""

import csv
import io
import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.export_service import COLUMNS
from tests.fixtures import NUCLEI_JSONL, SAMPLE_REPORTS, ZAP_JSON

TRACKER_CSV = SAMPLE_REPORTS / "tracker" / "findings-tracker.csv"
ZAP_PDF = SAMPLE_REPORTS / "zap" / "zap-automation.pdf"
HEADERS = [header for header, _, _ in COLUMNS]


def upload(client, project_id, path=None, *, name=None, content=None, headers=None, expect=201):
    body = content if content is not None else path.read_bytes()
    response = client.post("/api/scans/upload", data={"project_id": str(project_id)},
                           files={"file": (name or path.name, body, "application/octet-stream")}, headers=headers or {})
    assert response.status_code == expect, response.text
    return response.json()


@pytest.fixture
def alice(login_as):
    """A signed-in Security Analyst; `alice.user` is her account."""
    return login_as("Security Analyst", name="Alice", email="alice@example.com")


@pytest.fixture
def loaded(client, make_project, alice):
    project = make_project(name="Juice Shop")
    upload(client, project["id"], ZAP_JSON)
    upload(client, project["id"], NUCLEI_JSONL)
    return project


def export(client, fmt, headers=None, **params):
    response = client.get("/api/exports/findings", params={"format": fmt, **params}, headers=headers or {})
    assert response.status_code == 200, response.text
    return response


# --- Findings exports ----------------------------------------------------------------------

def test_csv_export_matches_the_database(client, loaded, sqlite_rows):
    response = export(client, "csv")
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0] == HEADERS
    assert len(rows) - 1 == sqlite_rows("SELECT COUNT(*) FROM findings")[0][0] == 48
    by_id = {int(r[0]): dict(zip(HEADERS, r)) for r in rows[1:]}
    for fid, title, evidence, severity, key in sqlite_rows("SELECT id, normalized_title, evidence, normalized_severity, finding_key FROM findings"):
        row = by_id[fid]
        assert (row["Title"], row["Evidence"], row["Severity"], row["Fingerprint"]) == (title, evidence or "", severity or "", key)


def test_exports_apply_the_same_filters_as_the_findings_list(client, loaded):
    listed = client.get("/api/findings", params={"severity": "Medium", "source_tool": "OWASP ZAP"}).json()["total"]
    data = json.loads(export(client, "json", severity="Medium", source_tool="OWASP ZAP").content)
    assert data["count"] == listed == len(data["findings"]) == 9
    assert data["filters"] == {"severities": ["Medium"], "source_tools": ["OWASP ZAP"]}
    assert {f["severity"] for f in data["findings"]} == {"Medium"}


def test_xlsx_export_opens_and_stores_text_not_formulas(client, make_project):
    project = make_project(name="Injection")
    csv_text = 'Title,Severity,Evidence\n"=cmd|\' /C calc\'!A0",High,=HYPERLINK("http://evil")\n+SUM(1),Low,@x\n'
    upload(client, project["id"], name="inj.csv", content=csv_text.encode())

    workbook = load_workbook(io.BytesIO(export(client, "xlsx").content))
    sheet = workbook["Findings"]
    assert [c.value for c in sheet[1]] == HEADERS
    values = {sheet.cell(row=r, column=3).value: sheet.cell(row=r, column=21) for r in range(2, sheet.max_row + 1)}
    assert set(values) == {"=cmd|' /C calc'!A0", "+SUM(1)"}
    assert all(cell.data_type == "s" for cell in values.values())  # stored as text, never evaluated
    assert {cell.value for cell in values.values()} == {'=HYPERLINK("http://evil")', "@x"}
    assert workbook["About"]["B2"].value == 2

    rows = list(csv.reader(io.StringIO(export(client, "csv").content.decode("utf-8-sig"))))
    titles = {r[2] for r in rows[1:]}
    assert titles == {"'=cmd|' /C calc'!A0", "'+SUM(1)"}  # CSV: neutralised with a leading quote

    json_titles = {f["title"] for f in json.loads(export(client, "json").content)["findings"]}
    assert json_titles == {"=cmd|' /C calc'!A0", "+SUM(1)"}  # JSON keeps the exact value


def test_exports_reflect_current_workflow_state(client, loaded, alice):
    fid = client.get("/api/findings", params={"search": "swagger"}).json()["items"][0]["id"]
    alice.put(f"/api/findings/{fid}", json={"status": "Accepted Risk", "assigned_to": alice.user["id"]})
    row = next(f for f in json.loads(export(client, "json").content)["findings"] if f["id"] == fid)
    assert (row["status"], row["assigned_to"]) == ("Accepted Risk", "Alice")


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
def test_exports_reimport_to_the_same_findings(client, loaded, make_project, sqlite_rows, fmt):
    content = export(client, fmt).content
    target = make_project(name=f"Reimport {fmt}")
    result = upload(client, target["id"], name=f"export.{fmt}", content=content)
    assert result["records_parsed"] == 48
    assert result["parser"] == f"Imported report {fmt.upper()}"
    original = sqlite_rows(
        "SELECT normalized_title, normalized_severity, COALESCE(evidence,''), COALESCE(endpoint,'') FROM findings WHERE project_id = ?",
        (loaded["id"],))
    reimported = sqlite_rows(
        "SELECT normalized_title, normalized_severity, COALESCE(evidence,''), COALESCE(endpoint,'') FROM findings WHERE project_id = ?",
        (target["id"],))
    assert sorted(original) == sorted(reimported)


def test_exports_are_saved_listed_downloadable_and_audited(client, loaded, alice, sqlite_rows, settings):
    response = export(alice, "json")
    name = re.search(r'filename="([^"]+)"', response.headers["content-disposition"]).group(1)
    assert (settings.export_path / name).read_bytes() == response.content
    listing = client.get("/api/exports").json()
    assert listing[0]["name"] == name and listing[0]["kind"] == "findings" and listing[0]["format"] == "json"
    assert client.get(f"/api/exports/files/{name}").content == response.content
    user_id, details = sqlite_rows("SELECT user_id, details FROM audit_logs WHERE action = 'report.exported'")[0]
    assert user_id == alice.user["id"] and json.loads(details)["count"] == 48


@pytest.mark.parametrize("name", ["..%2F..%2Fdata%2Fvapt.db", "../vapt.db", ".hidden", "missing.csv"])
def test_saved_export_names_cannot_escape_the_export_directory(client, name):
    assert client.get(f"/api/exports/files/{name}").status_code in (400, 404)


# --- Project report ------------------------------------------------------------------------

def test_project_report_contains_every_section_from_current_data(client, loaded, sqlite_rows):
    response = client.get(f"/api/exports/projects/{loaded['id']}/report")
    assert response.status_code == 200 and response.headers["content-type"].startswith("text/html")
    page = response.text
    for section in ("Project information", "Executive summary", "Risk summary", "Severity distribution",
                    "Medium findings (9)", "Low findings (17)", "Critical findings (0)", "High findings (0)",
                    "Detailed findings", "Remediation summary", "Scan information", "Findings register"):
        assert section in page, section
    total = sqlite_rows("SELECT COUNT(*) FROM findings WHERE project_id = ?", (loaded["id"],))[0][0]
    assert f"{total} unique finding(s) are recorded: 9 Medium, 17 Low, 22 Informational." in page
    assert "Content-Security-Policy" in page and "<script" not in page.lower().replace("&lt;script", "")
    assert "Not Provided" in page
    assert page.count('<article class="finding">') == total


def test_project_report_escapes_untrusted_scanner_text(client, make_project):
    project = make_project(name="Escaping")
    content = 'Title,Severity,Evidence,Description\n"<img src=x onerror=alert(1)>",High,<script>alert(document.cookie)</script>,a & b\n'
    upload(client, project["id"], name="xss.csv", content=content.encode())
    page = client.get(f"/api/exports/projects/{project['id']}/report").text
    assert "<script>alert" not in page and "<img src=x" not in page
    assert "&lt;script&gt;alert(document.cookie)&lt;/script&gt;" in page
    assert "&lt;img src=x onerror=alert(1)&gt;" in page and "a &amp; b" in page


def test_project_report_marks_not_extracted_fields(client, make_project):
    project = make_project(name="PDF")
    upload(client, project["id"], ZAP_PDF)
    page = client.get(f"/api/exports/projects/{project['id']}/report").text
    assert "Not Extracted" in page  # the ZAP PDF shows no Parameter row


def test_project_report_for_empty_and_missing_projects(client, make_project):
    empty = make_project(name="Empty")
    page = client.get(f"/api/exports/projects/{empty['id']}/report").text
    assert "No findings have been recorded for this project." in page and "No scans have been processed" in page
    assert client.get("/api/exports/projects/999/report").status_code == 404


# --- Bulk upload ---------------------------------------------------------------------------

def bulk(client, project_id, files, headers=None, **form):
    data = {"project_id": str(project_id), **{k: str(v).lower() for k, v in form.items()}}
    return client.post("/api/scans/upload-bulk", data=data, files=[("files", f) for f in files], headers=headers or {})


def test_bulk_upload_with_mixed_files(client, make_project, sqlite_rows, settings, alice):
    project = make_project(name="Bulk")
    files = [
        ("zap.json", ZAP_JSON.read_bytes(), "application/json"),
        ("nuclei.jsonl", NUCLEI_JSONL.read_bytes(), "application/octet-stream"),
        ("zap.pdf", ZAP_PDF.read_bytes(), "application/pdf"),
        ("broken.json", b'{"site": [', "application/json"),
        ("tool.exe", b"MZ\x90\x00", "application/octet-stream"),
        ("zap-again.json", ZAP_JSON.read_bytes(), "application/json"),
    ]
    response = bulk(alice, project["id"], files)
    assert response.status_code == 200, response.text
    result = response.json()
    assert [(r["filename"], r["outcome"]) for r in result["results"]] == [
        ("zap.json", "created"), ("nuclei.jsonl", "created"), ("zap.pdf", "created"),
        ("broken.json", "failed"), ("tool.exe", "rejected"), ("zap-again.json", "duplicate"),
    ]
    assert (result["total"], result["succeeded"], result["failed"], result["duplicates"]) == (6, 3, 2, 1)
    db_findings = sqlite_rows("SELECT COUNT(*) FROM findings WHERE project_id = ?", (project["id"],))[0][0]
    assert result["findings_added"] == db_findings == sum(r["new_findings"] for r in result["results"])
    assert "not valid JSON" in result["results"][3]["message"] and "not supported" in result["results"][4]["message"]
    # Originals kept for every processed or failed file; nothing for rejected or skipped duplicates.
    assert sqlite_rows("SELECT processing_status, COUNT(*) FROM scans WHERE project_id = ? GROUP BY 1 ORDER BY 1", (project["id"],)) == [
        ("Completed", 3), ("Failed", 1)]
    assert len([p for p in settings.upload_path.rglob("*") if p.is_file()]) == 4
    audit = json.loads(sqlite_rows("SELECT details FROM audit_logs WHERE action = 'reports.bulk_uploaded'")[0][0])
    assert (audit["succeeded"], audit["failed"], audit["duplicates"]) == (3, 2, 1)


def test_bulk_upload_can_process_duplicates_when_asked(client, make_project):
    project = make_project(name="Dupes")
    files = [("a.json", ZAP_JSON.read_bytes(), "application/json"), ("b.json", ZAP_JSON.read_bytes(), "application/json")]
    result = bulk(client, project["id"], files, allow_duplicates=True).json()
    assert [r["outcome"] for r in result["results"]] == ["created", "created"]
    assert result["results"][1]["new_findings"] == 0  # same findings, matched by fingerprint


def test_bulk_upload_limits(client, make_project):
    project = make_project(name="Limits")
    too_many = [(f"r{i}.json", b"{}", "application/json") for i in range(26)]
    assert bulk(client, project["id"], too_many).status_code == 400
    assert client.post("/api/scans/upload-bulk", data={"project_id": str(project["id"])}).status_code == 422
    assert bulk(client, 999, [("a.json", ZAP_JSON.read_bytes(), "application/json")]).status_code == 404


# --- Upload validation for the new formats ---------------------------------------------------

@pytest.mark.parametrize(
    "name, content, content_type, status",
    [
        ("report.xlsx", b"not a zip", "application/octet-stream", 415),
        ("report.pdf", b"<html>not a pdf</html>", "application/pdf", 415),
        ("report.csv", b"MZ\x90\x00", "text/csv", 415),
        ("report.csv", b"Title,Severity\nA,Low\n", "application/vnd.ms-excel", 201),  # what browsers on Windows send
        ("report.xlsx", TRACKER_CSV.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 415),
    ],
    ids=["xlsx-not-zip", "pdf-not-pdf", "csv-binary", "csv-windows-content-type", "csv-renamed-xlsx"],
)
def test_new_formats_are_validated(client, make_project, name, content, content_type, status):
    project = make_project(name="Validation")
    response = client.post("/api/scans/upload", data={"project_id": str(project["id"])}, files={"file": (name, content, content_type)})
    assert response.status_code == status, response.text
