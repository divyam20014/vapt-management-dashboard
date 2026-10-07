"""Phase 7 hardening: pre-body upload checks, response headers, SQLite settings, failure handling."""

import asyncio
import io
import sqlite3

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.main import create_app
from tests.conftest import new_client
from tests.fixtures import NUCLEI_JSONL, ZAP_JSON


def upload(client, project_id, name, content, **data):
    return client.post(
        "/api/scans/upload",
        data={"project_id": str(project_id), **data},
        files={"file": (name, content, "application/octet-stream")},
    )


@pytest.fixture
def project(client):
    return client.post("/api/projects", json={"name": "Hardening"}).json()


# --- Upload guard: refuse before reading the body ----------------------------------------------

def _run_guard(app, headers: dict[str, str], path="/api/scans/upload"):
    """Drive the ASGI app directly and record how much of the body it reads."""
    reads = {"body_messages": 0}
    sent = []

    async def receive():
        reads["body_messages"] += 1
        return {"type": "http.request", "body": b"x" * 1024, "more_body": True}

    async def send(message):
        sent.append(message)

    scope = {
        "type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": "POST", "scheme": "http",
        "path": path, "raw_path": path.encode(), "query_string": b"", "root_path": "",
        "headers": [(k.lower().encode(), v.encode()) for k, v in headers.items()],
        "client": ("127.0.0.1", 1234), "server": ("testserver", 80),
    }
    asyncio.run(app(scope, receive, send))
    status = next(m["status"] for m in sent if m["type"] == "http.response.start")
    return status, reads["body_messages"]


def test_anonymous_and_unauthorised_uploads_are_refused_before_the_body_is_read(app, client, login_as):
    csrf = {"X-Requested-With": "XMLHttpRequest", "Content-Type": "multipart/form-data; boundary=x"}
    assert _run_guard(app, {**csrf, "Content-Length": "999999"}) == (401, 0)
    assert _run_guard(app, {**csrf, "Content-Length": "999999"}, "/api/scans/upload-bulk") == (401, 0)
    assert _run_guard(app, {**csrf, "Content-Length": "999999", "Cookie": "vapt_session=forged"}) == (401, 0)
    viewer = login_as("Viewer")
    cookie = f"vapt_session={viewer.cookies.get('vapt_session')}"
    assert _run_guard(app, {**csrf, "Content-Length": "999999", "Cookie": cookie}) == (403, 0)
    # A declared size above the limit is refused even before the session is looked up.
    assert _run_guard(app, {**csrf, "Content-Length": str(10 * 1024**3)}) == (413, 0)


def test_streamed_body_over_the_limit_is_stopped(settings, project, client):
    settings.max_upload_size_mb = 1

    def chunks():  # no Content-Length: the size is only known while streaming
        boundary = b"--b\r\nContent-Disposition: form-data; name=\"file\"; filename=\"big.json\"\r\n\r\n"
        yield boundary
        for _ in range(4):
            yield b" " * (1024 * 1024)

    response = client.post(
        "/api/scans/upload", content=chunks(), headers={"Content-Type": "multipart/form-data; boundary=b"}
    )
    assert response.status_code == 413
    assert response.json()["error"]["code"] == "file_too_large"


def test_uploads_still_work_for_permitted_users(client, project, login_as):
    analyst = login_as("Security Analyst")
    assert upload(analyst, project["id"], NUCLEI_JSONL.name, NUCLEI_JSONL.read_bytes()).status_code == 201


# --- Response headers -------------------------------------------------------------------------

def test_security_headers(client, tmp_path, settings):
    response = client.get("/api/projects")
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "DENY"
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["referrer-policy"] == "same-origin"

    dist = tmp_path / "dist"
    dist.mkdir()
    (dist / "index.html").write_text("<!doctype html><title>x</title>", encoding="utf-8")
    settings.frontend_dist = str(dist)
    with new_client(create_app(settings)) as spa:
        page = spa.get("/findings")
        assert page.status_code == 200
        csp = page.headers["content-security-policy"]
        assert "script-src 'self'" in csp and "frame-ancestors 'none'" in csp and "unsafe-eval" not in csp
        assert "content-security-policy" not in spa.get("/docs").headers  # FastAPI's docs need their CDN assets


# --- SQLite settings --------------------------------------------------------------------------

def test_sqlite_waits_for_locks_and_uses_wal(app, client):
    with app.state.engine.connect() as conn:
        assert conn.execute(text("PRAGMA busy_timeout")).scalar() == 30000
        assert conn.execute(text("PRAGMA journal_mode")).scalar().lower() == "wal"
        assert conn.execute(text("PRAGMA foreign_keys")).scalar() == 1


def test_wal_can_be_turned_off(settings):
    settings.sqlite_wal = False
    with new_client(create_app(settings)):
        pass
    with sqlite3.connect(settings.sqlite_path) as conn:
        assert conn.execute("PRAGMA journal_mode").fetchone()[0].lower() != "wal"


# --- Failure handling -------------------------------------------------------------------------

def test_a_failed_upload_can_be_retried(client, project):
    broken = b'{"site": ['
    assert upload(client, project["id"], "zap.json", broken).status_code == 422
    retry = upload(client, project["id"], "zap.json", broken)
    assert retry.status_code == 422  # processed again, not reported as "already uploaded"
    assert upload(client, project["id"], ZAP_JSON.name, ZAP_JSON.read_bytes()).status_code == 201
    assert upload(client, project["id"], ZAP_JSON.name, ZAP_JSON.read_bytes()).status_code == 409


def test_unexpected_errors_while_saving_roll_back_and_keep_the_original(client, project, monkeypatch, settings, sqlite_rows):
    def explode(*args, **kwargs):
        raise RuntimeError("unexpected")

    monkeypatch.setattr("app.services.duplicate_service.detect_potential_duplicates", explode)
    response = upload(client, project["id"], ZAP_JSON.name, ZAP_JSON.read_bytes())
    assert response.status_code == 500
    assert response.json()["error"]["code"] == "processing_failed"
    scan_id = response.json()["error"]["details"]["scan_id"]
    assert sqlite_rows("SELECT processing_status FROM scans WHERE id = ?", (scan_id,)) == [("Failed",)]
    assert sqlite_rows("SELECT COUNT(*) FROM findings") == [(0,)] and sqlite_rows("SELECT COUNT(*) FROM finding_occurrences") == [(0,)]
    stored = sqlite_rows("SELECT stored_file_path FROM scans WHERE id = ?", (scan_id,))[0][0]
    assert (settings.upload_path / stored).read_bytes() == ZAP_JSON.read_bytes()
    assert "unexpected" not in response.text  # internal details are not returned


# --- Search, lists and exports ------------------------------------------------------------------

def test_search_treats_percent_and_underscore_literally(client, project):
    upload(client, project["id"], ZAP_JSON.name, ZAP_JSON.read_bytes())
    total = client.get("/api/findings").json()["total"]
    assert client.get("/api/findings", params={"search": "%"}).json()["total"] == 0
    assert client.get("/api/findings", params={"search": "_"}).json()["total"] < total
    assert client.get("/api/projects", params={"search": "%"}).json()["total"] == 0
    assert client.get("/api/projects", params={"search": "Harden"}).json()["total"] == 1


def test_bulk_upload_stores_every_record_as_an_occurrence(client, project, sqlite_rows):
    """The bulk INSERT path keeps one verbatim occurrence per report record."""
    result = upload(client, project["id"], ZAP_JSON.name, ZAP_JSON.read_bytes()).json()
    occurrences = sqlite_rows("SELECT COUNT(*) FROM finding_occurrences")[0][0]
    assert occurrences == 45  # ZAP baseline instances
    assert sqlite_rows("SELECT COUNT(*) FROM finding_status_history WHERE old_status IS NULL")[0][0] == result["new_findings"]
    assert sqlite_rows("SELECT COUNT(*) FROM finding_occurrences WHERE raw_record IS NULL") == [(0,)]


def test_xlsx_export_streams_text_cells(client, project):
    upload(client, project["id"], ZAP_JSON.name, ZAP_JSON.read_bytes())
    content = client.get("/api/exports/findings", params={"format": "xlsx"}).content
    sheet = load_workbook(io.BytesIO(content))["Findings"]
    assert sheet["A1"].value == "ID" and sheet["A1"].font.bold
    assert sheet.freeze_panes == "A2"
    assert all(c.data_type == "s" for row in sheet.iter_rows(min_row=2) for c in row if isinstance(c.value, str))
