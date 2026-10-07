"""Build the tracker-style fixtures (CSV, XLSX, JSON) from the real scanner outputs in this folder.

Every finding value is copied verbatim from zap/zap-baseline.json and nuclei/nuclei.jsonl; only the
layout (column names, a title row, an empty Status column) imitates a typical hand-maintained VAPT
tracker. No finding, severity or value is invented. Re-run with:

    python sample-reports/tools/make_tracker_fixtures.py
"""

import csv
import json
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tracker"

# Typical tracker headings (deliberately varied) -> source value
HEADERS = ["Sr. No.", "Vulnerability", "Risk Rating", "Affected URL", "HTTP Method", "Parameter", "Description",
           "Proof of Concept", "Remediation", "CWE ID", "Tool", "Status", "Tester Notes"]


def zap_rows() -> list[list[str]]:
    report = json.loads((ROOT / "zap" / "zap-baseline.json").read_text(encoding="utf-8"))
    rows = []
    for site in report["site"]:
        for alert in site["alerts"]:
            risk = alert["riskdesc"].split(" (")[0]
            cwe = alert["cweid"] if alert["cweid"] not in ("-1", "0") else ""
            for instance in alert["instances"]:
                rows.append([alert["alert"], risk, instance["uri"], instance["method"], instance["param"],
                             alert["desc"], instance["evidence"], alert["solution"], cwe, "OWASP ZAP"])
    return rows


def nuclei_rows() -> list[list[str]]:
    rows = []
    for line in (ROOT / "nuclei" / "nuclei.jsonl").read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        result = json.loads(line)
        info = result["info"]
        cwe = ", ".join(info.get("classification", {}).get("cwe-id") or [])
        name = info["name"] + (f": {result['matcher-name']}" if result.get("matcher-name") else "")
        rows.append([name, info["severity"], result["matched-at"], "", "", info.get("description", "").strip(),
                     "\n".join(result.get("extracted-results") or []), info.get("remediation", ""), cwe, "Nuclei"])
    return rows


def main() -> None:
    OUT.mkdir(exist_ok=True)
    data = zap_rows() + nuclei_rows()
    table = [[str(i), *row, "", ""] for i, row in enumerate(data, start=1)]  # empty Status / Tester Notes

    with open(OUT / "findings-tracker.csv", "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(table)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Juice Shop assessment - findings tracker (fixture)"])
    summary.append(["Generated from sample-reports/zap/zap-baseline.json and sample-reports/nuclei/nuclei.jsonl"])
    sheet = workbook.create_sheet("Findings")
    sheet.append(["Juice Shop assessment - findings"])  # a title row above the headings, as trackers often have
    sheet.append([])
    sheet.append(HEADERS)
    for row in table:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"  # store as text, never as a formula
    workbook.save(OUT / "findings-tracker.xlsx")

    keys = ["issue", "severity", "url", "method", "param", "details", "evidence", "solution", "cwe", "scanner"]
    items = [dict(zip(keys, row)) for row in data]
    (OUT / "findings.json").write_text(json.dumps({"vulnerabilities": items}, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"{len(data)} rows written to {OUT}")


if __name__ == "__main__":
    main()
